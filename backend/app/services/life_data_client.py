"""
抖音来客后台数据客户端

用于获取来客后台（life-data.cn）的直播数据和门店/视频数据
采用方案C：模拟来客后台API请求

创建日期: 2026-01-13
更新日期: 2026-01-31 - 添加门店概览导出功能获取视频数和门店评分
"""

import httpx
import logging
import asyncio
from typing import Optional, Dict, Any, List
from datetime import datetime, timedelta
from io import BytesIO

logger = logging.getLogger(__name__)


class LifeDataError(Exception):
    """来客后台数据抓取基础异常"""


class LifeDataAuthError(LifeDataError):
    """来客后台认证失效异常"""


class LifeDataPartialFetchError(LifeDataError):
    """来客后台分页抓取不完整异常"""


class LifeDataClient:
    """
    抖音来客后台数据客户端

    用于获取：
    1. 直播数据（直播时长、成交金额等）
    2. 门店/视频数据（视频发布数、核销金额等）
    """

    # API基础URL
    BASE_URL = "https://www.life-data.cn"

    # 直播数据API
    LIVE_DATA_API = "/api/lowcode_api/query"
    LIVE_DATA_PARAMS = {
        "projectKey": "shengyijing",
        "pageId": "15024",
        "interface_id": "3a3a25a6"
    }

    # 门店/视频数据API
    POI_DATA_API = "/api/dito/query"

    def __init__(
        self,
        cookie: str,
        life_account_id: str,
        csrf_token: Optional[str] = None
    ):
        """
        初始化来客数据客户端

        Args:
            cookie: 来客后台Cookie字符串
            life_account_id: 来客商户账户ID
            csrf_token: x-secsdk-csrf-token (可选，如需要)
        """
        self.cookie = cookie
        self.life_account_id = life_account_id
        self.csrf_token = csrf_token
        self._client: Optional[httpx.AsyncClient] = None

    def _get_common_headers(self) -> Dict[str, str]:
        """获取公共请求头"""
        headers = {
            "accept": "application/json, text/plain, */*",
            "accept-language": "zh-CN,zh;q=0.9,en;q=0.8",
            "content-type": "application/json",
            "cookie": self.cookie,
            "life-account-id": self.life_account_id,
            "root-life-account-id": self.life_account_id,
            "origin": self.BASE_URL,
            "user-agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/143.0.0.0 Safari/537.36",
            "sec-ch-ua": '"Microsoft Edge";v="143", "Chromium";v="143", "Not A(Brand";v="24"',
            "sec-ch-ua-mobile": "?0",
            "sec-ch-ua-platform": '"Windows"',
            "sec-fetch-dest": "empty",
            "sec-fetch-mode": "cors",
            "sec-fetch-site": "same-origin",
        }

        if self.csrf_token:
            headers["x-secsdk-csrf-token"] = self.csrf_token

        return headers

    async def _get_client(self) -> httpx.AsyncClient:
        """获取HTTP客户端（懒加载）"""
        if self._client is None or self._client.is_closed:
            self._client = httpx.AsyncClient(
                timeout=httpx.Timeout(60.0),  # 导出可能较慢
                follow_redirects=True
            )
        return self._client

    async def close(self):
        """关闭HTTP客户端"""
        if self._client and not self._client.is_closed:
            await self._client.aclose()
            self._client = None

    async def export_live_data(
        self,
        room_type_filter: str = "ALL",
        indicators: Optional[List[str]] = None
    ) -> Dict[str, Any]:
        """
        导出直播数据

        Args:
            room_type_filter: 直播类型过滤 ("ALL" | "1" 商家自播 | "2" 达人一带一 | "3" 达人一带多)
            indicators: 需要获取的指标列表，默认获取全部

        Returns:
            包含导出任务信息和Excel下载URL的字典
        """
        if indicators is None:
            indicators = [
                "task_id", "room_id", "room_title", "unique_id",
                "room_type", "room_cover_url", "live_start_ts",
                "live_start_str", "duration", "room_share_qr_code_url",
                "room_type_tag", "gmv", "room_pay_cert_num_td",
                "live_watch_uv_td", "room_gpm"
            ]

        payload = {
            "interface_id": self.LIVE_DATA_PARAMS["interface_id"],
            "data_query": {
                "data_query_type": "data_query",
                "filters": {
                    "operator": "and",
                    "filters": [
                        {
                            "field": {"name": "room_type_filter"},
                            "filter_operator": "equal",
                            "value": room_type_filter
                        }
                    ]
                },
                "group_by": [],
                "indicators": indicators,
                "extra": {}
            }
        }

        url = f"{self.BASE_URL}{self.LIVE_DATA_API}"
        params = self.LIVE_DATA_PARAMS.copy()

        headers = self._get_common_headers()
        headers["referer"] = f"{self.BASE_URL}/live/my/chain/calendar"

        client = await self._get_client()

        try:
            response = await client.post(url, params=params, json=payload, headers=headers)
            response.raise_for_status()
            result = response.json()

            if result.get("code") != 0:
                logger.error(f"导出直播数据失败: {result.get('msg')}")
                raise Exception(f"API返回错误: {result.get('msg')}")

            return result
        except httpx.HTTPStatusError as e:
            logger.error(f"HTTP请求失败: {e}")
            raise
        except Exception as e:
            logger.error(f"导出直播数据异常: {e}")
            raise

    async def export_poi_data(
        self,
        start_date: Optional[str] = None,
        end_date: Optional[str] = None,
        date_type: str = "last_seven_days",
        indicators: Optional[List[str]] = None,
        download: bool = False,
        offset: int = 0,
        limit: int = 100
    ) -> Dict[str, Any]:
        """
        导出/查询门店数据

        Args:
            start_date: 开始日期 (格式: YYYY-MM-DD)
            end_date: 结束日期 (格式: YYYY-MM-DD)
            date_type: 日期类型 ("last_seven_days" | "last_thirty_days" | "custom" | "nature_month")
            indicators: 需要获取的指标列表，如 ["video_cnt_1d", "poi_score"]
            download: 是否触发异步导出（True=导出Excel，False=直接分页查询）
            offset: 数据偏移量（从第几条开始，0起始）
            limit: 每次获取的数据条数（当前来客接口已验证 400 可用）

        Returns:
            包含数据的字典
        """
        user_provided_range = start_date is not None or end_date is not None

        # 默认时间范围：最近7天
        if end_date is None:
            end_date = datetime.now().strftime("%Y-%m-%d")
        if start_date is None:
            start_date = (datetime.now() - timedelta(days=6)).strftime("%Y-%m-%d")
        if user_provided_range and date_type not in ("custom", "nature_month", "last_thirty_days"):
            date_type = "custom"

        # 构建 AllPoiList 模块参数
        # poi_sizer 需要使用完整的结构，包含 dimension、isSelectAll、structValue、type
        poi_sizer_struct = {
            "dimension": 101,
            "isSelectAll": True,
            "structValue": '{"ExpandToPoiAccount":true,"SearchAllAccountPoiType":6,"SearchAllAccountPoiStatus":0,"Selections":[],"RelationTypes":[1,12,5],"PermissionKeyList":["hermes.data.shengyijing_data_view"]}',
            "type": "struct"
        }

        all_poi_list_params = {
            "poi_id": [],
            "brand_id": [],
            "poi_type": [],
            "poi_sizer": poi_sizer_struct
        }

        # 如果是下载模式，添加download参数
        if download:
            all_poi_list_params["download"] = 1

        # 如果指定了indicators，添加到参数中
        if indicators:
            all_poi_list_params["indicators"] = indicators

        # 分页参数（使用offset/limit，API不支持page/page_size）
        if not download:
            all_poi_list_params["offset"] = offset
            all_poi_list_params["limit"] = limit

        payload = {
            "biz_params": {
                "path": "/store/my/chain/poi/overview",
                "first_render": False,
                "common_params": {
                    "end_date": end_date,
                    "date_type": date_type,
                    "start_date": start_date
                },
                "module_params": {
                    "AllPoiList": all_poi_list_params
                }
            }
        }

        url = f"{self.BASE_URL}{self.POI_DATA_API}"

        headers = self._get_common_headers()
        headers["referer"] = f"{self.BASE_URL}/store/my/chain/poi/overview"

        client = await self._get_client()

        try:
            response = await client.post(url, json=payload, headers=headers)
            response.raise_for_status()

            # 调试：检查响应内容
            content_type = response.headers.get("content-type", "")
            response_text = response.text
            logger.warning(f"export_poi_data响应: status={response.status_code}, content-type={content_type}, 长度={len(response_text)}, 前200字符={response_text[:200]}")

            if not response_text or response_text.strip() == '':
                raise Exception("API返回空响应")

            result = response.json()

            result_code = result.get("code")
            result_msg = result.get("message") or result.get("msg") or "未知错误"
            if result_code == 401 or "请登录后再访问" in result_msg:
                logger.error(f"来客后台认证失效: {result_msg}")
                raise LifeDataAuthError(f"来客后台认证失效: {result_msg}")
            if result_code != 0:
                logger.error(f"导出门店数据失败: code={result_code}, message={result_msg}")
                raise LifeDataError(f"API返回错误(code={result_code}): {result_msg}")

            return result
        except httpx.HTTPStatusError as e:
            logger.error(f"HTTP请求失败: {e}")
            if e.response is not None and e.response.status_code == 401:
                raise LifeDataAuthError("来客后台认证失效: HTTP 401") from e
            raise
        except Exception as e:
            logger.error(f"导出门店数据异常: {e}")
            raise

    async def download_excel(self, download_url: str) -> bytes:
        """
        下载Excel文件

        Args:
            download_url: Excel文件下载URL

        Returns:
            Excel文件的字节内容
        """
        client = await self._get_client()

        try:
            response = await client.get(download_url, headers=self._get_common_headers())
            response.raise_for_status()
            return response.content
        except Exception as e:
            logger.error(f"下载Excel失败: {e}")
            raise

    async def _poll_export_task(self, task_id: str, start_date: str, end_date: str, max_retries: int = 30, interval: float = 2.0) -> Optional[str]:
        """
        轮询导出任务状态，等待Excel生成完成

        Args:
            task_id: 导出任务ID
            start_date: 开始日期
            end_date: 结束日期
            max_retries: 最大重试次数
            interval: 轮询间隔（秒）

        Returns:
            Excel下载URL，失败返回None
        """
        url = f"{self.BASE_URL}{self.POI_DATA_API}"
        headers = self._get_common_headers()
        headers["referer"] = f"{self.BASE_URL}/store/my/chain/list"
        client = await self._get_client()

        # 构建完整的轮询payload，与原始请求保持一致的结构
        poi_sizer_struct = {
            "dimension": 101,
            "isSelectAll": True,
            "structValue": '{"ExpandToPoiAccount":true,"SearchAllAccountPoiType":6,"SearchAllAccountPoiStatus":0,"Selections":[],"RelationTypes":[1,12,5],"PermissionKeyList":["hermes.data.shengyijing_data_view"]}',
            "type": "struct"
        }

        for i in range(max_retries):
            await asyncio.sleep(interval)
            try:
                # 查询导出任务状态的payload - 需要包含完整结构
                poll_payload = {
                    "biz_params": {
                        "path": "/store/my/rank",
                        "first_render": False,
                        "common_params": {
                            "end_date": end_date,
                            "date_type": "custom",
                            "start_date": start_date
                        },
                        "module_params": {
                            "AllPoiList": {
                                "poi_id": [],
                                "brand_id": [],
                                "poi_type": [],
                                "poi_sizer": poi_sizer_struct,
                                "download": 1,
                                "indicators": ["verify_amount_1d", "verify_cert_cnt_1d", "video_cnt_1d", "poi_score"],
                                "task_id": int(task_id)
                            }
                        }
                    }
                }
                response = await client.post(url, json=poll_payload, headers=headers)
                response.raise_for_status()

                # 调试：检查响应内容
                response_text = response.text
                if not response_text or response_text.strip() == '':
                    logger.warning(f"轮询({i+1}): 响应体为空, status_code={response.status_code}")
                    continue

                result = response.json()

                if result.get("code") == 0:
                    data = result.get("data", [])
                    # 调试日志：每次打印轮询结果
                    if isinstance(data, list) and data:
                        first_item = data[0] if isinstance(data[0], dict) else {}
                        logger.warning(f"轮询({i+1}): data=list, status={first_item.get('status')}, url存在={bool(first_item.get('url'))}, task_id={first_item.get('task_id')}")
                    elif isinstance(data, dict):
                        logger.warning(f"轮询({i+1}): data=dict, keys={list(data.keys())[:10]}")
                    else:
                        logger.warning(f"轮询({i+1}): data类型={type(data).__name__}, 值={str(data)[:200]}")
                    # 处理data是列表的情况
                    if isinstance(data, list) and data:
                        first_item = data[0] if isinstance(data[0], dict) else {}
                        download_url = first_item.get("url")
                        status = first_item.get("status")
                        if download_url and status == "success":
                            logger.info(f"导出任务完成，下载URL: {download_url[:100]}...")
                            return download_url
                        if status == "failed":
                            logger.error("导出任务失败")
                            return None
                    elif isinstance(data, dict):
                        module_data = data.get("AllPoiList", data)
                        download_url = module_data.get("download_url") or module_data.get("url")
                        if download_url:
                            logger.info(f"导出任务完成，下载URL: {download_url}")
                            return download_url
                        status = module_data.get("status", "")
                        if status == "failed":
                            logger.error("导出任务失败")
                            return None
                logger.debug(f"导出任务轮询中... ({i+1}/{max_retries})")
            except Exception as e:
                logger.warning(f"轮询导出任务异常: {e}")

        logger.error(f"导出任务超时（{max_retries}次轮询）")
        return None

    @staticmethod
    def _extract_poi_ranking_from_layout(response_data) -> tuple:
        """
        从嵌套的layout结构中递归查找poiRanking数据

        来客后台API (dito/query) 返回的数据是页面布局结构，
        门店数据嵌套在 layout -> children -> ... -> data.poiRanking 中。

        Args:
            response_data: API响应的完整data字段

        Returns:
            (store_list, total) 门店数据列表和总数
        """
        def _search(obj):
            if isinstance(obj, dict):
                # 检查当前dict是否包含poiRanking
                if "poiRanking" in obj:
                    poi_ranking = obj["poiRanking"]
                    if isinstance(poi_ranking, dict):
                        data_list = poi_ranking.get("data", [])
                        total = poi_ranking.get("total", 0)
                        if isinstance(data_list, list):
                            return data_list, total
                # 递归搜索所有值
                for value in obj.values():
                    result = _search(value)
                    if result:
                        return result
            elif isinstance(obj, list):
                for item in obj:
                    result = _search(item)
                    if result:
                        return result
            return None

        result = _search(response_data)
        if result:
            return result
        return [], 0

    def _extract_poi_page(self, result: Dict[str, Any], offset: int, log_label: str) -> tuple[List[Dict[str, Any]], int]:
        """从API响应中提取分页门店数据和总数"""
        data_field = result.get("data", {})
        total = result.get("total")

        if isinstance(data_field, dict):
            store_list, nested_total = self._extract_poi_ranking_from_layout(data_field)
            if nested_total:
                total = nested_total
            if total is None:
                total = len(store_list)
            logger.warning(
                f"{log_label}API响应(offset={offset}): "
                f"从layout结构提取到 {len(store_list)} 条门店数据, 总计 {total} 条"
            )
            return store_list, int(total or 0)

        if isinstance(data_field, list):
            store_list = data_field
            if total is None:
                total = len(store_list)
            logger.warning(
                f"{log_label}API响应(offset={offset}): "
                f"直接列表格式 {len(store_list)} 条, 总计 {total} 条"
            )
            return store_list, int(total or 0)

        logger.warning(
            f"{log_label}offset={offset}返回未知格式: type={type(data_field).__name__}"
        )
        return [], int(total or 0)

    async def _fetch_paginated_poi_rows(
        self,
        *,
        start_date: str,
        end_date: str,
        date_type: str,
        indicators: List[str],
        log_label: str,
    ) -> tuple[List[Dict[str, Any]], int, int]:
        """统一处理门店分页抓取，并在分页不完整时抛出异常"""
        offset = 0
        # 来客接口支持 400 条/页。用 400 可以把全量门店分页请求从约 50 次降到约 13 次，
        # 降低短时间高频请求触发空响应的概率。
        limit = 400
        total_fetched = 0
        expected_total: Optional[int] = None
        page_num = 0
        all_rows: List[Dict[str, Any]] = []
        max_retries = 4

        while True:
            page_num += 1
            store_list: List[Dict[str, Any]] = []
            page_total: Optional[int] = expected_total
            last_error: Optional[Exception] = None

            for retry in range(max_retries):
                try:
                    result = await self.export_poi_data(
                        start_date=start_date,
                        end_date=end_date,
                        date_type=date_type,
                        indicators=indicators,
                        download=False,
                        offset=offset,
                        limit=limit
                    )
                    store_list, page_total = self._extract_poi_page(result, offset, log_label)

                    if page_total is not None and expected_total is None:
                        expected_total = page_total
                    elif page_total not in (None, 0) and expected_total not in (None, 0) and page_total != expected_total:
                        logger.warning(
                            f"{log_label}总数发生变化: 初始={expected_total}, 当前页返回={page_total}, offset={offset}"
                        )

                    if store_list:
                        break

                    if (page_total or 0) == 0 and total_fetched == 0:
                        logger.warning(f"{log_label}分页查询无数据")
                        return [], 0, page_num

                    last_error = LifeDataPartialFetchError(
                        f"{log_label}offset={offset} 返回空数据，已获取 {total_fetched}/{expected_total or page_total or '?'} 条"
                    )
                except LifeDataAuthError:
                    logger.error(f"{log_label}offset={offset} 认证失效")
                    raise
                except Exception as e:
                    last_error = e
                    logger.warning(f"{log_label}offset={offset} 请求异常(第{retry + 1}次): {e}")

                if retry < max_retries - 1:
                    retry_delay = getattr(self, "_poi_page_retry_delay_seconds", 30)
                    if retry_delay > 0:
                        await asyncio.sleep(retry_delay * (retry + 1))  # 递增等待，给来客接口限流窗口恢复时间

            if last_error and not store_list:
                raise LifeDataPartialFetchError(
                    f"{log_label}分页抓取中断，offset={offset}，已获取 {total_fetched}/{expected_total or page_total or '?'} 条"
                ) from last_error

            if page_total is not None and expected_total is None:
                expected_total = page_total

            all_rows.extend(store_list)
            total_fetched += len(store_list)
            offset += len(store_list)

            if expected_total is not None and total_fetched >= expected_total:
                break

            # 每页成功后延迟，避免触发来客接口短时间分页限流
            page_delay = getattr(self, "_poi_page_success_delay_seconds", 8)
            if page_delay > 0:
                await asyncio.sleep(page_delay)

            if len(store_list) < limit:
                raise LifeDataPartialFetchError(
                    f"{log_label}分页提前结束，offset={offset - len(store_list)} 仅返回 {len(store_list)} 条，"
                    f"已获取 {total_fetched}/{expected_total or '?'} 条"
                )

            if offset > 5000:
                raise LifeDataPartialFetchError(
                    f"{log_label}分页超过安全上限5000，已获取 {total_fetched}/{expected_total or '?'} 条"
                )

        if expected_total is not None and total_fetched < expected_total:
            raise LifeDataPartialFetchError(
                f"{log_label}分页抓取不完整，实际获取 {total_fetched} 条，预期 {expected_total} 条"
            )

        return all_rows, expected_total or total_fetched, page_num

    async def _probe_poi_total(
        self,
        *,
        start_date: str,
        end_date: str,
        date_type: str,
        indicators: List[str],
        log_label: str,
    ) -> Optional[int]:
        """轻量探测门店总数，用于校验Excel导出结果是否完整。"""
        cache_key = (start_date, end_date, date_type, tuple(indicators))
        if not hasattr(self, "_poi_total_probe_cache"):
            self._poi_total_probe_cache = {}

        if cache_key in self._poi_total_probe_cache:
            return self._poi_total_probe_cache[cache_key]

        last_error: Optional[Exception] = None
        max_retries = 3

        for retry in range(max_retries):
            try:
                result = await self.export_poi_data(
                    start_date=start_date,
                    end_date=end_date,
                    date_type=date_type,
                    indicators=indicators,
                    download=False,
                    offset=0,
                    limit=1
                )
                _, total = self._extract_poi_page(result, 0, f"{log_label}总数探测")
                total_int = int(total or 0)
                self._poi_total_probe_cache[cache_key] = total_int
                return total_int
            except LifeDataAuthError:
                raise
            except Exception as e:
                last_error = e
                logger.warning(
                    f"{log_label}总数探测异常(第{retry + 1}次): {e}"
                )
                if retry < max_retries - 1:
                    await asyncio.sleep(2)

        logger.warning(
            f"{log_label}总数探测失败({last_error})，跳过Excel完整性校验"
        )
        return None

    async def _validate_excel_result_count(
        self,
        *,
        start_date: str,
        end_date: str,
        indicators: List[str],
        actual_count: int,
        log_label: str,
    ) -> None:
        """校验Excel导出结果是否完整，不完整时抛出异常阻止部分写库。"""
        expected_total = await self._probe_poi_total(
            start_date=start_date,
            end_date=end_date,
            date_type="custom",
            indicators=indicators,
            log_label=log_label,
        )

        if expected_total is None:
            return

        if actual_count < expected_total:
            raise LifeDataPartialFetchError(
                f"{log_label}结果不完整，实际 {actual_count} 条，预期 {expected_total} 条"
            )

        if actual_count > expected_total:
            logger.warning(
                f"{log_label}结果条数({actual_count})大于探测总数({expected_total})，继续使用导出结果"
            )

    @staticmethod
    def _parse_poi_excel_bytes(excel_bytes: bytes) -> List[Dict[str, Any]]:
        """解析Excel字节数据，返回门店数据行列表（列名使用中文标题）"""
        import openpyxl
        from io import BytesIO as _BytesIO
        wb = openpyxl.load_workbook(_BytesIO(excel_bytes), read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        header = next(rows_iter, None)
        if not header:
            wb.close()
            return []
        col_map: Dict[str, int] = {}
        for i, v in enumerate(header):
            if v is not None:
                title = str(v).strip()
                if title not in col_map:
                    col_map[title] = i
        result = []
        for row in rows_iter:
            if not row or all(v is None for v in row):
                continue
            item: Dict[str, Any] = {}
            for title, idx in col_map.items():
                if idx < len(row):
                    item[title] = row[idx]
            result.append(item)
        wb.close()
        return result

    async def fetch_all_poi_data_excel(
        self,
        start_date: Optional[str] = None,
        end_date: Optional[str] = None
    ) -> Dict[str, Dict[str, Any]]:
        """
        通过Excel下载方式一次性获取指定日期范围内所有门店数据

        使用 /store/my/rank 路径 + download:1 触发Excel异步导出，
        下载后解析Excel，返回按poi_id索引的门店数据字典。

        Returns:
            {poi_id: {"poi_name": str, "verify_amount": str, "verify_cert_cnt": int,
                      "video_cnt_1d": int, "poi_score": float}, ...}
        """
        if end_date is None:
            end_date = (datetime.now() - timedelta(days=1)).strftime("%Y-%m-%d")
        if start_date is None:
            start_date = (datetime.now() - timedelta(days=30)).strftime("%Y-%m-%d")

        cache_key = f"{start_date}_{end_date}"
        if hasattr(self, "_cached_all_poi_excel") and getattr(self, "_cached_all_poi_excel_key", "") == cache_key:
            logger.warning(f"复用缓存的Excel导出数据({start_date}~{end_date})，共 {len(self._cached_all_poi_excel)} 个门店")
            return self._cached_all_poi_excel

        indicators = ["verify_amount_1d", "verify_cert_cnt_1d", "video_cnt_1d", "poi_score"]
        import time as _time
        task_id = int(_time.time() * 1000)

        poi_sizer_struct = {
            "dimension": 101,
            "isSelectAll": True,
            "structValue": '{"ExpandToPoiAccount":true,"SearchAllAccountPoiType":6,"SearchAllAccountPoiStatus":0,"Selections":[],"RelationTypes":[1,12,5],"PermissionKeyList":["hermes.data.shengyijing_data_view"]}',
            "type": "struct"
        }

        payload = {
            "biz_params": {
                "path": "/store/my/rank",
                "first_render": False,
                "common_params": {
                    "end_date": end_date,
                    "date_type": "custom",
                    "start_date": start_date
                },
                "module_params": {
                    "AllPoiList": {
                        "poi_id": [],
                        "brand_id": [],
                        "poi_type": [],
                        "poi_sizer": poi_sizer_struct,
                        "download": 1,
                        "indicators": indicators,
                        "task_id": task_id
                    }
                }
            }
        }

        url = f"{self.BASE_URL}{self.POI_DATA_API}"
        headers = self._get_common_headers()
        headers["referer"] = f"{self.BASE_URL}/store/my/chain/list"
        client_http = await self._get_client()

        logger.warning(f"fetch_all_poi_data_excel: 发送导出请求, date={start_date}~{end_date}, task_id={task_id}")
        response = await client_http.post(url, json=payload, headers=headers)
        response.raise_for_status()

        resp_text = response.text
        logger.warning(f"fetch_all_poi_data_excel: 响应长度={len(resp_text)}, 前300字符={resp_text[:300]}")
        if not resp_text or resp_text.strip() == '':
            raise LifeDataError("Excel导出触发：API返回空响应，可能Cookie已失效")

        result = response.json()
        result_code = result.get("code")
        result_msg = result.get("msg") or result.get("message") or "未知错误"
        if result_code == 401 or "请登录" in str(result_msg):
            raise LifeDataAuthError(f"来客后台认证失效: {result_msg}")
        if result_code != 0:
            raise LifeDataError(f"触发Excel导出失败: code={result_code}, msg={result_msg}")

        data = result.get("data", [])
        download_url = None
        returned_task_id = None

        if isinstance(data, list) and data:
            first_item = data[0] if isinstance(data[0], dict) else {}
            status = first_item.get("status", "")
            url_val = first_item.get("url", "")
            returned_task_id = first_item.get("task_id")
            logger.warning(f"fetch_all_poi_data_excel: 初始响应 status={status}, url存在={bool(url_val)}, task_id={returned_task_id}")
            if status == "success" and url_val:
                download_url = url_val
                logger.warning(f"Excel导出任务立即完成，URL已获取")
            elif status in ("submit_error", "failed"):
                extra = first_item.get("extra", {})
                raise LifeDataError(f"Excel导出任务提交失败: status={status}, extra={extra}")

        if not download_url and returned_task_id:
            logger.warning(f"Excel导出任务异步等待，task_id={returned_task_id}，开始轮询...")
            download_url = await self._poll_export_task(str(returned_task_id), start_date, end_date)

        if not download_url:
            raise LifeDataError("无法获取Excel下载URL，导出任务可能失败")

        excel_bytes = await self.download_excel(download_url)
        rows = self._parse_poi_excel_bytes(excel_bytes)
        logger.warning(f"Excel解析完成，共 {len(rows)} 行数据")

        result_dict: Dict[str, Dict[str, Any]] = {}
        for row in rows:
            poi_id_val = row.get("门店ID")
            poi_id = str(poi_id_val or "").strip()
            if not poi_id or poi_id in ("None", "null", ""):
                continue
            poi_name = str(row.get("门店名称") or "").strip()
            raw_amount = self._safe_int_static(row.get("门店核销金额", 0))
            result_dict[poi_id] = {
                "poi_name": poi_name,
                "verify_amount": self._format_amount_fen_to_yuan(raw_amount),
                "verify_cert_cnt": self._safe_int_static(row.get("门店核销券数", 0)),
                "video_cnt_1d": self._safe_int_static(row.get("新发布门店关联视频数", 0)),
                "poi_score": self._safe_float_static(row.get("门店评分", 0)),
            }

        await self._validate_excel_result_count(
            start_date=start_date,
            end_date=end_date,
            indicators=indicators,
            actual_count=len(result_dict),
            log_label="Excel导出",
        )

        logger.warning(f"fetch_all_poi_data_excel: 解析到 {len(result_dict)} 个有效门店")
        self._cached_all_poi_excel = result_dict
        self._cached_all_poi_excel_key = cache_key
        return result_dict

    async def _fetch_combined_poi_data(
        self,
        start_date: Optional[str] = None,
        end_date: Optional[str] = None
    ) -> List[Dict[str, Any]]:
        """
        一次性获取所有门店的全部指标（概览+核销），避免重复分页请求。
        结果缓存在 self._cached_combined_rows 中供后续调用复用。
        """
        if start_date is None:
            today = datetime.now()
            start_date = today.replace(day=1).strftime("%Y-%m-%d")
        if end_date is None:
            end_date = (datetime.now() - timedelta(days=1)).strftime("%Y-%m-%d")

        cache_key = f"{start_date}_{end_date}"
        if hasattr(self, '_cached_combined_rows') and getattr(self, '_cached_combined_key', '') == cache_key:
            logger.warning(f"复用缓存的合并分页数据，共 {len(self._cached_combined_rows)} 条")
            return self._cached_combined_rows

        # 方案1: 优先 Excel 导出
        try:
            all_data = await self.fetch_all_poi_data_excel(start_date, end_date)
            rows = []
            for poi_id, data in all_data.items():
                rows.append({
                    "poi_id_str": poi_id,
                    "poi_name": data.get("poi_name", ""),
                    "video_cnt_1d": data.get("video_cnt_1d", 0),
                    "poi_score": data.get("poi_score", 0),
                    "verify_amount_1d": 0,
                    "verify_cert_cnt_1d": data.get("verify_cert_cnt", 0),
                    "_verify_amount_formatted": data.get("verify_amount", "¥0.00"),
                })
            self._cached_combined_rows = rows
            self._cached_combined_key = cache_key
            logger.warning(f"Excel导出成功，获取到 {len(rows)} 条合并数据")
            return rows
        except Exception as excel_err:
            logger.warning(f"Excel导出失败({excel_err})，回退到合并分页查询...")

        # 方案2: 一次分页获取4个指标
        store_rows, total, page_num = await self._fetch_paginated_poi_rows(
            start_date=start_date,
            end_date=end_date,
            date_type="nature_month",
            indicators=["verify_amount_1d", "verify_cert_cnt_1d", "video_cnt_1d", "poi_score"],
            log_label="合并数据"
        )
        self._cached_combined_rows = store_rows
        self._cached_combined_key = cache_key
        logger.warning(f"合并分页完成，获取到 {len(store_rows)} 条 (共{page_num}页, 总计{total}条)")
        return store_rows

    async def fetch_poi_overview_data(
        self,
        start_date: Optional[str] = None,
        end_date: Optional[str] = None
    ) -> Dict[str, Dict[str, Any]]:
        """
        获取门店概览数据（video_cnt_1d 和 poi_score）
        日期范围：当月1号到昨天
        """
        if start_date is None:
            today = datetime.now()
            start_date = today.replace(day=1).strftime("%Y-%m-%d")
        if end_date is None:
            end_date = (datetime.now() - timedelta(days=1)).strftime("%Y-%m-%d")

        try:
            # 优先 Excel
            try:
                all_data = await self.fetch_all_poi_data_excel(start_date, end_date)
                poi_data = {}
                for poi_id, data in all_data.items():
                    poi_data[poi_id] = {
                        "video_cnt_1d": data.get("video_cnt_1d", 0),
                        "poi_score": data.get("poi_score", 0.0),
                        "poi_name": data.get("poi_name", ""),
                    }
                logger.warning(f"从Excel获取到 {len(poi_data)} 个门店概览数据")
                return poi_data
            except Exception as excel_err:
                logger.warning(f"Excel导出失败({excel_err})，回退到分页查询...")

            # 回退分页
            store_rows, total, page_num = await self._fetch_paginated_poi_rows(
                start_date=start_date,
                end_date=end_date,
                date_type="nature_month",
                indicators=["video_cnt_1d", "poi_score"],
                log_label="门店概览"
            )
            poi_data = {}
            for item in store_rows:
                if not isinstance(item, dict):
                    continue
                poi_id = str(item.get("poi_id_str", "") or item.get("poi_id", "") or "")
                if poi_id and poi_id not in ("None", "null", ""):
                    poi_data[poi_id] = {
                        "video_cnt_1d": self._safe_int_static(item.get("video_cnt_1d", 0)),
                        "poi_score": self._safe_float_static(item.get("poi_score", 0)),
                        "poi_name": item.get("poi_name", ""),
                    }
            logger.warning(f"从分页获取到 {len(poi_data)} 个门店概览数据 (共{page_num}页)")
            return poi_data
        except Exception as e:
            logger.error(f"获取门店概览数据失败: {e}")
            import traceback
            logger.error(traceback.format_exc())
            raise

    async def fetch_poi_verify_data(
        self,
        start_date: Optional[str] = None,
        end_date: Optional[str] = None,
        date_type: str = "last_thirty_days",
        log_label: str = "核销数据"
    ) -> Dict[str, Dict[str, Any]]:
        """
        获取门店核销数据（verify_amount 和 verify_cert_cnt）

        默认取近30天；传入 date_type="custom" 且指定 start_date/end_date 时，
        可获取当月实时日期范围（例如 5月1~14号）的核销数据。
        """
        if start_date is None:
            if date_type == "last_thirty_days":
                start_date = (datetime.now() - timedelta(days=30)).strftime("%Y-%m-%d")
            else:
                today = datetime.now()
                start_date = today.replace(day=1).strftime("%Y-%m-%d")
        if end_date is None:
            end_date = (datetime.now() - timedelta(days=1)).strftime("%Y-%m-%d")

        try:
            # 优先 Excel
            try:
                all_data = await self.fetch_all_poi_data_excel(start_date, end_date)
                poi_verify = {}
                for poi_id, data in all_data.items():
                    poi_verify[poi_id] = {
                        "verify_amount": data.get("verify_amount", "¥0.00"),
                        "verify_cert_cnt": data.get("verify_cert_cnt", 0),
                        "poi_name": data.get("poi_name", ""),
                    }
                logger.warning(f"从Excel获取到 {len(poi_verify)} 个门店{log_label}")
                return poi_verify
            except Exception as excel_err:
                logger.warning(f"Excel导出失败({excel_err})，回退到分页查询...")

            # 回退分页
            store_rows, total, page_num = await self._fetch_paginated_poi_rows(
                start_date=start_date,
                end_date=end_date,
                date_type=date_type,
                indicators=["verify_amount_1d", "verify_cert_cnt_1d"],
                log_label=log_label
            )
            poi_verify = {}
            for item in store_rows:
                if not isinstance(item, dict):
                    continue
                poi_id = str(item.get("poi_id_str", "") or item.get("poi_id", "") or "")
                if poi_id and poi_id not in ("None", "null", ""):
                    raw_amount = self._safe_int_static(item.get("verify_amount_1d", 0))
                    poi_verify[poi_id] = {
                        "verify_amount": self._format_amount_fen_to_yuan(raw_amount),
                        "verify_cert_cnt": self._safe_int_static(item.get("verify_cert_cnt_1d", 0)),
                        "poi_name": item.get("poi_name", ""),
                    }
            logger.warning(f"从分页获取到 {len(poi_verify)} 个门店{log_label} (共{page_num}页)")
            return poi_verify
        except Exception as e:
            logger.error(f"获取门店核销数据失败: {e}")
            import traceback
            logger.error(traceback.format_exc())
            raise

    @staticmethod
    def _format_amount_fen_to_yuan(fen_value: int) -> str:
        """
        将分转换为元，返回带¥前缀的格式化字符串

        Args:
            fen_value: 金额（分）

        Returns:
            格式化字符串，如 "¥1,785.20"
        """
        yuan = fen_value / 100
        return f"¥{yuan:,.2f}"

    @staticmethod
    def _safe_int_static(val) -> int:
        """安全转换为整数"""
        if val is None:
            return 0
        if isinstance(val, (int, float)):
            return int(val)
        if isinstance(val, str):
            try:
                # 处理可能的逗号分隔数字
                return int(val.replace(",", "").replace("，", ""))
            except ValueError:
                return 0
        return 0

    @staticmethod
    def _safe_float_static(val) -> float:
        """安全转换为浮点数"""
        if val is None:
            return 0.0
        if isinstance(val, (int, float)):
            return float(val)
        if isinstance(val, str):
            try:
                return float(val.replace(",", "").replace("，", ""))
            except ValueError:
                return 0.0
        return 0.0

    async def get_live_data_parsed(
        self,
        room_type_filter: str = "ALL"
    ) -> Dict[str, Any]:
        """
        获取并解析直播数据

        Args:
            room_type_filter: 直播类型过滤

        Returns:
            解析后的直播数据，包含：
            - total: 总数
            - data: 直播列表（包含duration、gmv等字段）
            - excel_url: Excel下载链接
        """
        result = await self.export_live_data(room_type_filter)

        data_list = result.get("data", [])
        excel_url = None

        if data_list and len(data_list) > 0:
            excel_url = data_list[0].get("url")

        return {
            "total": result.get("total", 0),
            "meta": result.get("meta", []),
            "data": data_list,
            "excel_url": excel_url,
            "raw_response": result
        }

    async def get_poi_data_parsed(
        self,
        start_date: Optional[str] = None,
        end_date: Optional[str] = None
    ) -> Dict[str, Any]:
        """
        获取并解析门店/视频数据

        Args:
            start_date: 开始日期
            end_date: 结束日期

        Returns:
            解析后的门店数据，包含：
            - total: 总数
            - data: 门店列表（包含video_cnt_1d、verify_amount_1d等字段）
            - excel_url: Excel下载链接
        """
        result = await self.export_poi_data(start_date, end_date)

        data_list = result.get("data", [])
        excel_url = None

        if data_list and len(data_list) > 0:
            excel_url = data_list[0].get("url")

        return {
            "total": result.get("total", 0),
            "meta": result.get("meta", []),
            "data": data_list,
            "excel_url": excel_url,
            "raw_response": result
        }

    async def get_summary_data(
        self,
        start_date: Optional[str] = None,
        end_date: Optional[str] = None
    ) -> Dict[str, Any]:
        """
        获取汇总数据（直播时长、视频发布数、核销金额）

        这是一个便捷方法，同时获取直播和门店数据并提取关键指标

        Args:
            start_date: 开始日期
            end_date: 结束日期

        Returns:
            汇总数据字典，包含：
            - live_duration_total: 总直播时长（秒）
            - gmv_total: 总成交金额（分）
            - video_count_total: 视频发布总数
            - verify_amount_total: 核销总金额（分）
            - live_excel_url: 直播数据Excel下载链接
            - poi_excel_url: 门店数据Excel下载链接
        """
        # 并行获取两类数据
        import asyncio

        live_task = self.get_live_data_parsed()
        poi_task = self.get_poi_data_parsed(start_date, end_date)

        live_result, poi_result = await asyncio.gather(live_task, poi_task)

        # 提取关键指标
        summary = {
            "live_duration_total": 0,
            "gmv_total": 0,
            "video_count_total": 0,
            "verify_amount_total": 0,
            "live_excel_url": live_result.get("excel_url"),
            "poi_excel_url": poi_result.get("excel_url"),
            "query_date_range": {
                "start_date": start_date,
                "end_date": end_date
            }
        }

        # 从直播数据中提取
        for item in live_result.get("data", []):
            if item.get("duration"):
                summary["live_duration_total"] += int(item["duration"])
            if item.get("gmv"):
                summary["gmv_total"] += int(item["gmv"])

        # 从门店数据中提取
        for item in poi_result.get("data", []):
            if item.get("video_cnt_1d"):
                summary["video_count_total"] += int(item["video_cnt_1d"])
            if item.get("verify_amount_1d"):
                summary["verify_amount_total"] += int(item["verify_amount_1d"])

        return summary


def parse_excel_to_dict(excel_bytes: bytes) -> List[Dict[str, Any]]:
    """
    解析Excel文件为字典列表

    Args:
        excel_bytes: Excel文件字节内容

    Returns:
        包含每行数据的字典列表
    """
    try:
        import openpyxl
        from io import BytesIO

        workbook = openpyxl.load_workbook(BytesIO(excel_bytes), read_only=True)
        sheet = workbook.active

        rows = list(sheet.iter_rows(values_only=True))
        if len(rows) < 2:
            return []

        headers = rows[0]
        data = []

        for row in rows[1:]:
            item = {}
            for i, value in enumerate(row):
                if i < len(headers) and headers[i]:
                    item[str(headers[i])] = value
            data.append(item)

        return data
    except ImportError:
        logger.warning("openpyxl未安装，无法解析Excel文件")
        raise ImportError("需要安装openpyxl库: pip install openpyxl")
    except Exception as e:
        logger.error(f"解析Excel失败: {e}")
        raise
